from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any

try:
    import xlsxwriter
except ModuleNotFoundError:  # pragma: no cover - fallback keeps local dev working without network
    xlsxwriter = None
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from .models import Finding, SEVERITIES, STATUS_CLOSED, STATUS_OPEN

FINDING_COLUMNS = [
    ("record_no", "Kayıt No"),
    ("title", "Bulgu Başlığı"),
    ("severity", "Durum Seviyesi"),
    ("impact", "Bulgunun Etkisi"),
    ("description", "Bulgunun Açıklaması"),
    ("recommendation", "Çözüm Önerisi"),
    ("related_unit", "İlgili Birim / Kurum"),
    ("related_person", "İlgili Kişi"),
    ("status", "Durum"),
    ("due_date", "Termin Tarih"),
    ("new_due_date", "Yeni Termin"),
    ("note", "Not"),
    ("updated_at", "Son Güncelleme"),
    ("sla_status", "SLA Durumu"),
    ("active_due_date", "Aktif Termin"),
    ("delay_days", "Gecikme Günü"),
    ("last_action_note", "Son Aksiyon"),
]

HEADER_ALIASES = {
    "record_no": ["kayıt no", "kayit no", "no"],
    "title": ["bulgu adı", "bulgu adi", "bulgu başlığı", "bulgu basligi", "başlık", "baslik"],
    "severity": ["önem derecesi", "onem derecesi", "durum seviyesi", "seviye", "risk", "öncelik", "oncelik"],
    "impact": ["bulgunun etkisi", "etki"],
    "description": ["bulgunun açıklaması", "bulgunun aciklamasi", "açıklama", "aciklama"],
    "recommendation": ["çözüm önerisi", "cozum onerisi", "öneri", "oneri"],
    "related_unit": ["ilgili birim / kurum", "ilgili birim", "ilgili kurum", "ilgili"],
    "related_person": ["ilgili kişi", "ilgili kisi", "sorumlu", "sorumlu kisi"],
    "status": ["durum", "kapanma durumu"],
    "due_date": ["termin tarih", "termin tarihi", "termin"],
    "new_due_date": ["yeni termin", "yeni termin tarih", "yeni termin tarihi"],
    "note": ["not", "notlar", "açık not", "acik not"],
}

IMPORT_FIELDS = {
    "record_no",
    "title",
    "severity",
    "impact",
    "description",
    "recommendation",
    "related_unit",
    "related_person",
    "status",
    "due_date",
    "new_due_date",
    "note",
    "source",
}

NOTE_HEADER_ALIASES = {
    "Erişim Noktası": ["erişim noktası", "erisim noktasi"],
    "Kullanıcı Profili": ["kullanıcı profili", "kullanici profili"],
    "Bulgunun Tespit Edildiği Bileşen/Bileşenler": [
        "bulgunun tespit edildiği bileşen/bileşenler",
        "bulgunun tespit edildigi bilesen/bilesenler",
        "bulgunun tespit edildiği bileşen",
        "bulgunun tespit edildigi bilesen",
        "bileşen/bileşenler",
        "bilesen/bilesenler",
    ],
}

FORBIDDEN_IMPORT_FIELDS = {
    "active_due_date",
    "delay_days",
    "sla_status",
    "last_action_note",
    "is_overdue",
    "is_due_soon",
    "updated_at",
    "created_at",
    "id",
}


def normalize_header(value: Any) -> str:
    text = str(value or "").strip().lower().replace("i̇", "i")
    text = text.translate(str.maketrans({"ı": "i", "İ": "i", "ğ": "g", "ü": "u", "ş": "s", "ö": "o", "ç": "c"}))
    text = re.sub(r"[^a-z0-9/ ]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    return None


def normalize_status(value: Any) -> str:
    normalized = normalize_header(value)
    if normalized in {"kapatildi", "kapali", "closed", "tamamlandi"}:
        return STATUS_CLOSED
    return STATUS_OPEN


def normalize_severity(value: Any) -> str:
    normalized = normalize_header(value)
    for severity in SEVERITIES:
        if normalized == normalize_header(severity):
            return severity
    return "Orta"


def build_header_map(headers: list[Any]) -> dict[str, int]:
    normalized_headers = [normalize_header(header) for header in headers]
    header_map: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        alias_set = {normalize_header(alias) for alias in aliases}
        for index, normalized in enumerate(normalized_headers):
            if normalized in alias_set:
                header_map[field] = index
                break
    return header_map


def build_note_header_map(headers: list[Any]) -> dict[str, int]:
    normalized_headers = [normalize_header(header) for header in headers]
    header_map: dict[str, int] = {}
    for label, aliases in NOTE_HEADER_ALIASES.items():
        alias_set = {normalize_header(alias) for alias in aliases}
        for index, normalized in enumerate(normalized_headers):
            if normalized in alias_set:
                header_map[label] = index
                break
    return header_map


def find_header_row(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int], dict[str, int]]:
    best_row_index = 0
    best_header_map: dict[str, int] = {}
    best_note_header_map: dict[str, int] = {}
    best_score = -1
    for row_index, row in enumerate(rows):
        header_map = build_header_map(list(row))
        note_header_map = build_note_header_map(list(row))
        score = len(header_map) + len(note_header_map)
        if "title" in header_map:
            score += 10
        if score > best_score:
            best_row_index = row_index
            best_header_map = header_map
            best_note_header_map = note_header_map
            best_score = score
    return best_row_index, best_header_map, best_note_header_map


def allowed_import_fields() -> set[str]:
    model_fields = {column.name for column in Finding.__table__.columns}
    return (model_fields & IMPORT_FIELDS) - FORBIDDEN_IMPORT_FIELDS


def clean_import_payload(payload: dict[str, Any]) -> dict[str, Any]:
    allowed_fields = allowed_import_fields()
    return {field: value for field, value in payload.items() if field in allowed_fields}


def row_to_payload(row: tuple[Any, ...], header_map: dict[str, int], note_header_map: dict[str, int], fallback_no: int, source: str) -> dict[str, Any]:
    payload: dict[str, Any] = {field: "" for field in allowed_import_fields()}
    payload.update({"record_no": f"IMPORT-{fallback_no}", "severity": "Orta", "status": STATUS_OPEN, "due_date": None, "new_due_date": None, "source": source})
    for field, index in header_map.items():
        if field not in IMPORT_FIELDS:
            continue
        value = row[index] if index < len(row) else None
        if field in {"due_date", "new_due_date"}:
            payload[field] = parse_date(value)
        elif field == "status":
            payload[field] = normalize_status(value)
        elif field == "severity":
            payload[field] = normalize_severity(value)
        else:
            payload[field] = str(value or "").strip()
    extra_notes: list[str] = []
    for label, index in note_header_map.items():
        value = row[index] if index < len(row) else None
        text = str(value or "").strip()
        if text:
            extra_notes.append(f"{label}: {text}")
    if extra_notes:
        existing_note = str(payload.get("note") or "").strip()
        payload["note"] = "\n".join([part for part in [existing_note, *extra_notes] if part])
    return clean_import_payload(payload)


def read_excel(contents: bytes, filename: str = "Excel") -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    workbook = load_workbook(BytesIO(contents), data_only=True)
    if "Bulgular" not in workbook.sheetnames:
        raise ValueError('Excel içinde "Bulgular" sayfası bulunamadı')
    sheet = workbook["Bulgular"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    header_row_index, header_map, note_header_map = find_header_row(rows)
    errors: list[dict[str, Any]] = []
    payloads: list[dict[str, Any]] = []
    if "title" not in header_map:
        errors.append({"row": header_row_index + 1, "message": "Bulgu Adı kolonu bulunamadı"})
    for row_number, row in enumerate(rows[header_row_index + 1 :], start=header_row_index + 2):
        if not any(cell not in (None, "") for cell in row):
            continue
        payload = row_to_payload(row, header_map, note_header_map, row_number, filename)
        if not payload.get("title"):
            continue
        payloads.append(payload)
    return payloads, errors


def import_findings(db: Session, contents: bytes, filename: str, current_actor: str = "system", action_callback: Any | None = None) -> dict[str, Any]:
    payloads, errors = read_excel(contents, filename)
    created = 0
    updated = 0
    for payload in payloads:
        finding = db.query(Finding).filter(Finding.record_no == payload["record_no"]).one_or_none()
        if finding:
            for key, value in payload.items():
                setattr(finding, key, value)
            if action_callback:
                action_callback(db, finding, "imported_update", current_actor, note=f"Excel import güncellemesi: {filename}")
            updated += 1
        else:
            finding = Finding(**payload)
            db.add(finding)
            db.flush()
            if action_callback:
                action_callback(db, finding, "imported", current_actor, note=f"Excel import kaydı: {filename}")
            created += 1
    db.flush()
    return {"message": "İçe aktarma tamamlandı", "created": created, "updated": updated, "failed": len(errors), "errors": errors}


def build_import_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Bulgular"
    headers = [
        "Bulgu Adı",
        "Önem Derecesi",
        "Erişim Noktası",
        "Kullanıcı Profili",
        "Bulgunun Tespit Edildiği Bileşen/Bileşenler",
        "Bulgunun Etkisi",
        "Bulgunun Açıklaması",
        "Çözüm Önerisi",
        "İlgili Birim / Kurum",
        "İlgili Kişi",
        "Durum",
        "Termin Tarih",
        "Yeni Termin",
    ]
    sheet.append(["Sürat Kargo Sızma Testi-Güvenlik Bulguları - 2026"])
    sheet.append(headers)
    examples = [
        ["Örnek - Yönetim panelinde güçlü parola politikası eksik", "Orta", "Örnek URL", "Örnek kullanıcı", "Örnek web uygulaması", "Örnek etki açıklaması", "Bu satır örnek amaçlıdır; içe aktarmadan önce silinebilir.", "Örnek çözüm önerisi", "Örnek Birim", "Örnek Kişi", STATUS_OPEN, date.today(), None],
        ["Örnek - Güncel olmayan bileşen kullanımı", "Yüksek", "Örnek servis", "Örnek profil", "Örnek bileşen", "Örnek etki", "Bu kayıt gerçek sistem kaydı değildir.", "Örnek paket güncelleme aksiyonu", "Örnek BT", "Örnek Sorumlu", STATUS_OPEN, date.today(), date.today()],
        ["Örnek - Kapatılmış test bulgusu", "Düşük", "Örnek endpoint", "Örnek rol", "Örnek API", "Örnek düşük etki", "Şablon formatını göstermek için eklenmiştir.", "Örnek doğrulama", "Örnek Operasyon", "", STATUS_CLOSED, date.today(), None],
    ]
    for row in examples:
        sheet.append(row)
    title_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    sheet["A1"].font = Font(bold=True, size=13, color="0F2F57")
    sheet["A1"].fill = title_fill
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    for cell in sheet[2]:
        cell.fill = header_fill
        cell.font = header_font
    severity_validation = DataValidation(type="list", formula1='"Acil,Kritik,Yüksek,Orta,Düşük"', allow_blank=False)
    status_validation = DataValidation(type="list", formula1='"Devam Ediyor,Kapatıldı"', allow_blank=False)
    sheet.add_data_validation(severity_validation)
    sheet.add_data_validation(status_validation)
    severity_validation.add("B3:B500")
    status_validation.add("K3:K500")
    for row in range(3, 501):
        sheet[f"L{row}"].number_format = "dd.mm.yyyy"
        sheet[f"M{row}"].number_format = "dd.mm.yyyy"
    widths = [34, 18, 22, 22, 42, 32, 42, 34, 28, 24, 18, 16, 16]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=2, column=idx).column_letter].width = width
    sheet.freeze_panes = "A3"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def preview_findings(contents: bytes, filename: str) -> dict[str, Any]:
    payloads, errors = read_excel(contents, filename)
    return {"preview": payloads[:50], "total": len(payloads), "failed": len(errors), "errors": errors}


def excel_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value or ""


def finding_value(finding: Finding, field: str) -> Any:
    if field == "active_due_date":
        return finding.new_due_date or finding.due_date
    return getattr(finding, field, "")


def build_export(findings: list[Finding], summary: dict[str, Any]) -> bytes:
    if xlsxwriter is None:
        return build_export_openpyxl(findings, summary)
    output = BytesIO()
    workbook = xlsxwriter.Workbook(output, {"in_memory": True})
    summary_sheet = workbook.add_worksheet("Ozet")
    findings_sheet = workbook.add_worksheet("Bulgular")

    title_fmt = workbook.add_format({"bold": True, "font_size": 14, "font_color": "#0f2f57"})
    header_fmt = workbook.add_format({"bold": True, "bg_color": "#1F4E78", "font_color": "#FFFFFF", "border": 1})
    green_fmt = workbook.add_format({"bg_color": "#D9EAD3"})
    percent_fmt = workbook.add_format({"num_format": "0.00%"})

    summary_sheet.write(0, 0, "Sürat Kargo Sızma - Güvenlik Testi Bulgu Özeti", title_fmt)
    summary_sheet.write(2, 0, "Durum Seviyesi", header_fmt)
    summary_sheet.write(2, 1, "Toplam", header_fmt)
    summary_sheet.write(2, 2, "Kapatılan", header_fmt)
    summary_sheet.write(2, 3, "Açık Kalan", header_fmt)
    for row_idx, row in enumerate(summary["severity_rows"], start=3):
        summary_sheet.write(row_idx, 0, row["severity"])
        summary_sheet.write_number(row_idx, 1, row["total"])
        summary_sheet.write_number(row_idx, 2, row["closed"])
        summary_sheet.write_number(row_idx, 3, row["open"])
    footer = 4 + len(summary["severity_rows"])
    summary_sheet.write(footer, 0, "Kapanma Oranı")
    summary_sheet.write_number(footer, 1, (summary.get("closure_rate") or 0) / 100, percent_fmt)
    summary_sheet.write(footer + 1, 0, "Son Çalışma Saati")
    summary_sheet.write(footer + 1, 1, str(summary.get("last_work_time") or ""))
    summary_sheet.set_column("A:D", 24)

    for col, (_, header) in enumerate(FINDING_COLUMNS):
        findings_sheet.write(0, col, header, header_fmt)
    for row_idx, finding in enumerate(findings, start=1):
        is_closed = finding.status == STATUS_CLOSED
        for col, (field, _) in enumerate(FINDING_COLUMNS):
            fmt = green_fmt if is_closed and (col == 0 or 2 <= col <= 10) else None
            findings_sheet.write(row_idx, col, excel_value(finding_value(finding, field)), fmt)
    last_row = max(len(findings) + 1, 2)
    findings_sheet.data_validation(1, 8, last_row, 8, {"validate": "list", "source": [STATUS_OPEN, STATUS_CLOSED]})
    for idx, width in enumerate([16, 34, 16, 28, 42, 32, 28, 22, 16, 16, 16, 28, 22, 16, 16, 14, 24]):
        findings_sheet.set_column(idx, idx, width)
    findings_sheet.freeze_panes(1, 0)
    workbook.close()
    output.seek(0)
    return output.read()


def build_export_openpyxl(findings: list[Finding], summary: dict[str, Any]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Ozet"
    findings_sheet = workbook.create_sheet("Bulgular")
    green_fill = PatternFill("solid", fgColor="D9EAD3")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    summary_sheet.append(["Sürat Kargo Sızma - Güvenlik Testi Bulgu Özeti"])
    summary_sheet.append([])
    summary_sheet.append(["Durum Seviyesi", "Toplam", "Kapatılan", "Açık Kalan"])
    for cell in summary_sheet[3]:
        cell.fill = header_fill
        cell.font = header_font
    for row in summary["severity_rows"]:
        summary_sheet.append([row["severity"], row["total"], row["closed"], row["open"]])
    summary_sheet.append([])
    summary_sheet.append(["Kapanma Oranı", f"%{summary.get('closure_rate') or 0:.2f}"])
    summary_sheet.append(["Son Çalışma Saati", str(summary.get("last_work_time") or "")])
    findings_sheet.append([header for _, header in FINDING_COLUMNS])
    for cell in findings_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for finding in findings:
        findings_sheet.append([excel_value(finding_value(finding, field)) for field, _ in FINDING_COLUMNS])
        if finding.status == STATUS_CLOSED:
            for col_idx in [1, *range(3, 12)]:
                findings_sheet.cell(row=findings_sheet.max_row, column=col_idx).fill = green_fill
    validation = DataValidation(type="list", formula1='"Devam Ediyor,Kapatıldı"', allow_blank=False)
    findings_sheet.add_data_validation(validation)
    validation.add(f"I2:I{max(findings_sheet.max_row, 2)}")
    for sheet in (summary_sheet, findings_sheet):
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 60)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_defender_export(summary: dict[str, Any], vulnerabilities: list[dict[str, Any]], machines: list[dict[str, Any]], recommendations: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Defender_Ozet"
    vuln_sheet = workbook.create_sheet("Defender_CVE_Listesi")
    machine_sheet = workbook.create_sheet("Defender_Cihaz_Yazilim_CVE")
    rec_sheet = workbook.create_sheet("Defender_Oneriler")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    summary_rows = [
        ["Metrik", "Değer"],
        ["Toplam CVE", summary.get("total_cve", 0)],
        ["Kritik CVE", summary.get("critical_cve", 0)],
        ["Yüksek CVE", summary.get("high_cve", 0)],
        ["Etkilenen Cihaz", summary.get("affected_machines", 0)],
        ["Public Exploit", summary.get("public_exploit", 0)],
        ["Verified Exploit", summary.get("verified_exploit", 0)],
        ["Remediation Bekleyen", summary.get("pending_remediation", 0)],
        ["Son Sync", summary.get("last_sync", "")],
    ]
    for row in summary_rows:
        summary_sheet.append(row)

    def add_rows(sheet: Any, headers: list[str], rows: list[dict[str, Any]]) -> None:
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in rows:
            sheet.append([excel_value(row.get(header)) for header in headers])
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 60)

    add_rows(vuln_sheet, ["cve_id", "description", "severity", "cvss_v3", "exposed_machines", "public_exploit", "exploit_verified", "epss", "published_on", "updated_on", "status"], vulnerabilities)
    add_rows(machine_sheet, ["machine_name", "machine_id", "cve_id", "product_vendor", "product_name", "product_version", "severity", "fixing_kb_id", "recommendation_id", "remediation_status", "first_seen", "last_seen"], machines)
    add_rows(rec_sheet, ["recommendation_id", "recommendation_name", "product_name", "vendor", "recommendation_category", "severity_score", "exposed_machines", "remediation_type", "status", "exposure_impact"], recommendations)
    for cell in summary_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_defender_affected_devices_export(devices: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Etkilenen_Cihazlar"
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    columns = [
        ("Cihaz Adı", "device_name"),
        ("OS", "os_platform"),
        ("Vendor", "product_vendor"),
        ("Ürün / Uygulama", "product_name"),
        ("Ürün Versiyonu", "product_version"),
        ("CVE", "cve_id"),
        ("Seviye", "severity"),
        ("KB", "fixing_kb_id"),
        ("İlk Görülme", "first_seen"),
        ("Son Görülme", "last_seen"),
    ]
    sheet.append([label for label, _ in columns])
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for device in devices:
        sheet.append([excel_value(device.get(key)) for _, key in columns])
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 60)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def build_msrc_export(summary: dict[str, Any], vulnerabilities: list[dict[str, Any]]) -> bytes:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "MSRC_Ozet"
    vuln_sheet = workbook.create_sheet("MSRC_CVE_Listesi")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    summary_rows = [
        ["Metrik", "Değer"],
        ["Toplam CVE", summary.get("total_cve", 0)],
        ["Critical", summary.get("critical", 0)],
        ["High / Important", summary.get("high_important", 0)],
        ["Medium / Moderate", summary.get("moderate", 0)],
        ["Low", summary.get("low", 0)],
        ["Exploited", summary.get("exploited", 0)],
        ["Publicly Disclosed", summary.get("publicly_disclosed", 0)],
        ["KB Sayısı", summary.get("kb_count", 0)],
    ]
    for row in summary_rows:
        summary_sheet.append(row)
    for cell in summary_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    columns = [
        ("cve_id", "CVE"),
        ("title", "Başlık"),
        ("severity", "Seviye"),
        ("product", "Etkilenen Ürün / Uygulama"),
        ("kb_article", "KB"),
        ("impact", "Etki"),
        ("exploited", "İstismar Ediliyor mu?"),
        ("publicly_disclosed", "Kamuya Açık mı?"),
        ("release_month", "Yayın Ayı"),
        ("release_date", "Yayın Tarihi"),
    ]
    vuln_sheet.append([label for _, label in columns])
    for cell in vuln_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in vulnerabilities:
        values = []
        for key, _ in columns:
            value = row.get(key)
            if key in {"exploited", "publicly_disclosed"}:
                value = "Evet" if value else "Hayır"
            values.append(excel_value(value))
        vuln_sheet.append(values)

    for sheet in (summary_sheet, vuln_sheet):
        for column_cells in sheet.columns:
            max_length = max(len(str(cell.value or "")) for cell in column_cells)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 2, 12), 70)
        sheet.freeze_panes = "A2"
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
